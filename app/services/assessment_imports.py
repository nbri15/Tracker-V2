"""Parse, validate and execute reviewed assessment imports atomically."""
from __future__ import annotations
import base64
import csv
import io
import re
import zipfile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from collections import Counter
from openpyxl import load_workbook
from app.extensions import db
from app.models import AcademicYear, Pupil, SchoolClass, SubjectResult
from .assessments import CsvImportError, AssessmentValidationError, compute_subject_result_values, get_class_pupil_query, get_subject_setting
from .assessment_reliability import apply_result, cohort_year_group, setting_for_result

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


def read_upload(file, suffix):
    if not file or not file.filename or not file.filename.lower().endswith(suffix):
        raise CsvImportError(f'Choose a {suffix} file first.')
    data = file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise CsvImportError('The file is too large. Upload a file smaller than 5 MB.')
    return data


def parse_csv_bytes(data):
    try:
        text = data.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text), strict=True)
        headers = reader.fieldnames
        if not headers or any(not value or not value.strip() for value in headers) or len(set(headers)) != len(headers):
            raise CsvImportError('Use a header row with unique, named columns.')
        rows = []
        for index, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise CsvImportError(f'Row {index} has the wrong number of cells. Use the downloaded template.')
            if any(value.strip() for value in row.values()):
                rows.append(row)
            if len(rows) > MAX_ROWS:
                raise CsvImportError('Upload no more than 5,000 rows at a time.')
        if not rows:
            raise CsvImportError('The file has no pupil rows.')
        return rows
    except (UnicodeDecodeError, csv.Error) as exc:
        raise CsvImportError('The CSV could not be read. Save it as UTF-8 CSV using the template.') from exc


def parse_score(value, label):
    text = str(value if value is not None else '').strip()
    if not text:
        return None
    if not re.fullmatch(r'[+-]?\d+(?:\.0+)?', text):
        raise CsvImportError(f'{label} must be a whole number.')
    return int(text.split('.')[0])


def class_pupils(payload):
    school_class = db.session.get(SchoolClass, payload['class_id'])
    if not school_class or school_class.school_id != payload['school_id']:
        raise CsvImportError('The selected class is not available in this school.')
    return school_class, get_class_pupil_query(school_class, payload['academic_year']).filter(Pupil.is_active.is_(True)).all()


def process_subject(payload):
    school_class, pupils = class_pupils(payload)
    by_id = {pupil.id: pupil for pupil in pupils}
    by_name = {}
    for pupil in pupils:
        by_name.setdefault(pupil.full_name.casefold(), []).append(pupil)
    setting = get_subject_setting(school_class.year_group, payload['subject'], payload['term'], payload['academic_year'], school_id=payload['school_id'])
    rows = payload['rows']
    if not {'pupil_id', 'paper_1_score', 'paper_2_score'}.issubset(rows[0]) and not {'pupil', 'paper_1_score', 'paper_2_score'}.issubset(rows[0]):
        raise CsvImportError('Use columns pupil_id (or pupil), paper_1_score and paper_2_score. Download this assessment’s template.')
    existing = {row.pupil_id: row for row in SubjectResult.query.filter_by(school_id=payload['school_id'], academic_year=payload['academic_year'], subject=payload['subject'], term=payload['term']).filter(SubjectResult.pupil_id.in_(by_id)).all()}
    report = {'errors': [], 'rows': [], 'matched': 0, 'added': 0, 'updated': 0, 'blank': 0, 'total': len(pupils)}
    seen = set()
    for index, row in enumerate(rows, 2):
        pupil = None
        label = row.get('pupil') or row.get('pupil_id') or f'Row {index}'
        try:
            if str(row.get('pupil_id', '')).strip():
                raw_id = str(row['pupil_id']).strip()
                pupil = by_id.get(int(raw_id)) if raw_id.isdigit() else None
                if not pupil:
                    raise CsvImportError('Pupil ID is not in the selected class and school. Check the template.')
                if row.get('pupil') and row['pupil'].strip().casefold() != pupil.full_name.casefold():
                    raise CsvImportError('Pupil name and pupil ID do not match.')
            else:
                candidates = by_name.get(str(row.get('pupil', '')).strip().casefold(), [])
                if len(candidates) != 1:
                    raise CsvImportError('Name could not be matched uniquely. Use the pupil ID from the template.')
                pupil = candidates[0]
            label = pupil.full_name
            if pupil.id in seen:
                raise CsvImportError('This pupil appears more than once. Remove the duplicate row.')
            seen.add(pupil.id)
            if row.get('class_name') and row['class_name'].strip() != school_class.name:
                raise CsvImportError('Class name does not match the selected assessment.')
            if row.get('academic_year') and row['academic_year'].strip() != payload['academic_year']:
                raise CsvImportError('Academic year does not match the selected assessment.')
            result = existing.get(pupil.id)
            row_setting = setting_for_result(result, setting) if result else setting
            first = parse_score(row.get('paper_1_score'), row_setting.paper_1_name)
            second = parse_score(row.get('paper_2_score'), row_setting.paper_2_name)
            report['matched'] += 1
            if first is None and second is None:
                report['blank'] += 1
                report['rows'].append({'pupil': label, 'detail': 'Both scores blank: existing result will be kept.', 'status': 'warning'})
                continue
            if result:
                first = first if first is not None else result.paper_1_score
                second = second if second is not None else result.paper_2_score
            compute_subject_result_values(row_setting, first, second)
            if result and result.source in {'manual', 'gap'} and not payload.get('replace_existing'):
                raise CsvImportError('A manual/GAP result already exists. Tick “replace existing results” to review replacing it.')
            result = result or SubjectResult(pupil=pupil, pupil_id=pupil.id, school_id=payload['school_id'], academic_year=payload['academic_year'], subject=payload['subject'], term=payload['term'])
            apply_result(result, row_setting, first, second, assessment_year_group=result.assessment_year_group, source='csv')
            action = 'updated' if pupil.id in existing else 'added'
            report[action] += 1
            report['rows'].append({'pupil': label, 'detail': f"{action.title()}: {first if first is not None else 'blank'} + {second if second is not None else 'blank'}; {result.combined_percent if result.combined_percent is not None else 'incomplete'}; {result.band_label or 'no band'}", 'status': 'warning' if first is None or second is None else 'ready'})
        except (CsvImportError, AssessmentValidationError) as exc:
            message = f'{label} — {exc}'
            report['errors'].append(message)
            report['rows'].append({'pupil': label, 'detail': message, 'status': 'error'})
    report['missing'] = [pupil.full_name for pupil in pupils if pupil.id not in seen]
    db.session.flush()
    report['without_score'] = sum(1 for pupil in pupils if not (result := SubjectResult.query.filter_by(pupil_id=pupil.id, academic_year=payload['academic_year'], subject=payload['subject'], term=payload['term']).first()) or result.combined_percent is None)
    return report


def validate_combined(payload):
    from .csv_tools import COMBINED_SUBJECT_SCORE_COLUMNS, _split_pupil_name, _parse_year_group
    rows = payload['rows']
    errors, details, seen = [], [], set()
    if not {'class_name', 'year_group'}.issubset(rows[0]) or not ('pupil' in rows[0] or {'first_name', 'last_name'}.issubset(rows[0])):
        raise CsvImportError('Use the combined template with pupil, class_name and year_group columns.')
    for index, row in enumerate(rows, 2):
        try:
            name, _, _ = _split_pupil_name(row)
            if name.lower().startswith('example'):
                continue
            group = _parse_year_group(row.get('year_group'))
            class_name = (row.get('class_name') or '').strip()
            if not class_name:
                raise CsvImportError('Class name is missing.')
            year = row.get('academic_year') or payload['academic_year']
            if year != payload['academic_year']:
                raise CsvImportError('Academic year differs from the selected year. Import each year separately.')
            row['academic_year'] = year
            key = (name.casefold(), class_name.casefold())
            if key in seen:
                raise CsvImportError('Duplicate pupil row. Keep one row per pupil.')
            seen.add(key)
            school_class = SchoolClass.query.filter_by(school_id=payload['school_id'], name=class_name).first()
            from .assessments import get_school_working_academic_year
            if school_class and school_class.year_group != group and year == get_school_working_academic_year(payload['school_id']).name:
                raise CsvImportError('Year group does not match this class. This import cannot promote classes.')
            raw_id = row.get('pupil_id', '').strip()
            if raw_id:
                pupil = Pupil.query.filter_by(id=int(raw_id), school_id=payload['school_id']).first() if raw_id.isdigit() else None
                if not pupil or pupil.full_name.casefold() != name.casefold() or (school_class and pupil.class_id != school_class.id):
                    raise CsvImportError('Pupil ID, name and class must match this school.')
            matches = Pupil.query.filter_by(school_id=payload['school_id'], class_id=school_class.id).all() if school_class else []
            candidates = [pupil for pupil in matches if pupil.full_name.casefold() == name.casefold()]
            if len(candidates) > 1 and not raw_id:
                raise CsvImportError('This name is ambiguous. Add pupil_id to identify the correct pupil.')
            details.append({'pupil': name, 'status': 'ready' if candidates else 'warning', 'detail': 'Existing pupil matched.' if candidates else 'Will create a new pupil (combined roster import). Check this name carefully.'})
            for subject, terms in COMBINED_SUBJECT_SCORE_COLUMNS.items():
                for term, (first_key, second_key) in terms.items():
                    first, second = parse_score(row.get(first_key), first_key), parse_score(row.get(second_key), second_key)
                    if first is not None or second is not None:
                        setting = get_subject_setting(group, subject, term, year, school_id=payload['school_id'])
                        if candidates:
                            saved = SubjectResult.query.filter_by(pupil_id=candidates[0].id, academic_year=year, subject=subject, term=term).first()
                            if saved:
                                setting = setting_for_result(saved, setting)
                        compute_subject_result_values(setting, first, second)
        except (ValueError, AssessmentValidationError) as exc:
            errors.append(f'Row {index} ({row.get("pupil", "pupil")}): {exc}')
    return errors, details


def read_workbook(encoded):
    data = base64.b64decode(encoded)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 30 * 1024 * 1024:
                raise CsvImportError('The workbook is too large when unpacked. Use a smaller workbook.')
        return load_workbook(io.BytesIO(data), data_only=True)
    except (ValueError, zipfile.BadZipFile, KeyError, OSError, TypeError, ParseError, InvalidFileException) as exc:
        raise CsvImportError('The workbook could not be read. Upload a valid .xlsx template.') from exc


def validate_workbook(wb, school_id, academic_year):
    errors = []
    supported = {'Pupils', 'Maths', 'Reading', 'SPaG', 'Phonics', 'SATs'}
    for ws in wb:
        if ws.title not in supported:
            if ws.title.lower() not in {'instructions', 'readme'}:
                header_index = next((index for index in range(1, min(ws.max_row, 5) + 1) if any(str(cell.value or '').strip().lower() == 'pupil_id' for cell in ws[index])), 1)
                headings = [str(cell.value or '').strip().lower() for cell in ws[header_index]]
                identity_fields = {'pupil_id', 'pupil', 'class', 'subject', 'term', 'test_name', 'tracking_point', 'academic_year', 'year group', 'year_group'}
                populated = any(value not in (None, '') for row in ws.iter_rows(min_row=header_index + 1, values_only=True) for key, value in zip(headings, row) if key not in identity_fields)
                if populated:
                    errors.append(f'{ws.title}: this populated sheet is not supported by this importer. Use its dedicated tracker.')
            continue
        header_row = next((index for index in range(1, min(ws.max_row, 5) + 1) if any(str(cell.value or '').strip().lower() in {'pupil', 'pupil_name', 'pupil_id'} for cell in ws[index])), None)
        if not header_row:
            errors.append(f'{ws.title}: pupil header is missing. Use the workbook template.')
            continue
        headers = [str(cell.value or '').strip().lower().replace(' ', '_') for cell in ws[header_row]]
        if len([value for value in headers if value]) != len(set(value for value in headers if value)):
            errors.append(f'{ws.title}: duplicate column headings.')
        if ws.title in {'Maths', 'Reading', 'SPaG'}:
            first_aliases = {'arithmetic', 'paper_1', 'spelling'}
            second_aliases = {'reasoning', 'paper_2', 'grammar'}
            if 'term' not in headers or not first_aliases.intersection(headers) or not second_aliases.intersection(headers):
                errors.append(f'{ws.title}: include Term and both paper-score columns from the template.')
        seen = set()
        count = 0
        for index, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            count += 1
            if count > MAX_ROWS:
                errors.append(f'{ws.title}: more than 5,000 rows. Split the workbook.')
                break
            row = dict(zip(headers, values))
            name = str(row.get('pupil') or row.get('pupil_name') or '').strip()
            pupil_id = row.get('pupil_id')
            pupil = None
            if pupil_id not in (None, ''):
                try:
                    pupil = Pupil.query.filter_by(id=parse_score(pupil_id, 'Pupil ID'), school_id=school_id).first()
                except CsvImportError:
                    pupil = None
                if not pupil or (name and name.casefold() != pupil.full_name.casefold()):
                    errors.append(f'{ws.title} row {index}: pupil ID and name do not match this school.')
                elif row.get('class'):
                    from app.models import PupilClassHistory
                    history = PupilClassHistory.query.filter_by(school_id=school_id, pupil_id=pupil.id, academic_year=academic_year).first()
                    expected_class = history.class_name if history else pupil.school_class.name if pupil.school_class else ''
                    if str(row['class']).strip().casefold() != expected_class.casefold():
                        errors.append(f'{ws.title} row {index}: class does not match this pupil in {academic_year}.')
            key = (str(pupil.id if pupil else name).casefold(), str(row.get('class') or '').casefold(), str(row.get('term') or row.get('test_name') or row.get('assessment_point') or row.get('exam_number') or '').casefold())
            if key in seen:
                errors.append(f'{ws.title} row {index}: duplicate result for {name}.')
            seen.add(key)
            for field, value in row.items():
                if field in {'pupil_id', 'pupil', 'pupil_name', 'class', 'term', 'notes', 'gender', 'pp', 'send', 'laps', 'service_child', 'year_group', 'service', 'test_name', 'date', 'assessment_point', 'exam_number', 'writing_band'} or not field:
                    continue
                if value not in (None, ''):
                    try:
                        score = parse_score(value, field)
                        if score < 0:
                            raise CsvImportError('Scores cannot be negative.')
                        if 'scaled' in field and not 80 <= score <= 120:
                            raise CsvImportError('SATs scaled scores must be between 80 and 120.')
                    except CsvImportError as exc:
                        errors.append(f'{ws.title} row {index}, {name}: {exc}')
    return errors


def process_import(kind, payload):
    if kind == 'subject':
        return process_subject(payload)
    if kind == 'workbook':
        from app.admin.routes import _process_full_workbook
        wb = read_workbook(payload['workbook'])
        errors = validate_workbook(wb, payload['school_id'], payload['academic_year'])
        if errors:
            return {'errors': errors, 'rows': [], 'checked': 0}
        sats = process_sats_workbook(wb['SATs'],payload) if 'SATs' in wb.sheetnames else {'errors':[], 'rows':[], 'added':0, 'changed_results':0}
        if 'SATs' in wb.sheetnames:
            wb.remove(wb['SATs'])
        report = _process_full_workbook(wb, payload['academic_year'])
        report['errors'].extend(sats['errors'])
        report['rows'].extend(sats['rows'])
        report['added'] = report.get('added',0) + sats['added']
        report['changed_results'] = report.get('changed_results',0) + sats['changed_results']
        return report
    from .csv_tools import import_combined_results, import_reception_tracker, import_sats_tracker_results
    handlers = {'combined': import_combined_results, 'reception': import_reception_tracker, 'sats_tracker': import_sats_tracker_results}
    if kind not in handlers:
        raise CsvImportError('Choose a supported import type.')
    details = []
    if kind == 'combined':
        errors, details = validate_combined(payload)
        if errors:
            return {'errors': errors, 'rows': details, 'checked': len(payload['rows'])}
    else:
        seen = set()
        for row in payload['rows']:
            identity = tuple(row.get(key) for key in ('pupil_first_name', 'pupil_last_name', 'class_name', 'academic_year', 'exam_number', 'tracking_point'))
            if identity in seen:
                raise CsvImportError('Duplicate pupil/assessment rows. Remove duplicates before uploading.')
            seen.add(identity)
    summary = handlers[kind](payload['rows'])
    return {'matched': summary.pupils_matched, 'errors': summary.errors[:100] if summary.validation_errors else [], 'warnings': summary.errors if not summary.validation_errors else [], 'rows': details, 'checked': summary.rows_processed, 'created': summary.pupils_created, 'updated': summary.pupils_updated, 'added': summary.subject_results_created + summary.writing_results_created + summary.tracker_entries_created, 'changed_results': summary.subject_results_updated + summary.writing_results_updated + summary.tracker_entries_updated, 'skipped': summary.rows_skipped + summary.manual_results_skipped}


def simulate_import(kind, payload):
    """Validate every path with the actual writer, then roll all domain writes back."""
    transaction = db.session.begin_nested()
    try:
        return process_import(kind, payload)
    finally:
        transaction.rollback()


def save_sats_result(pupil, academic_year, exam_number, values):
    """Keep SATs' three maths papers/scaled scores separate from term percentages."""
    from app.models import SimpleSatsSetting, SatsResult
    if exam_number not in range(1, 5):
        raise CsvImportError('Choose a SATs assessment point from 1 to 4.')
    if cohort_year_group(pupil, academic_year) != 6:
        raise CsvImportError('SATs results require a Year 6 pupil in the selected year.')
    configured = SimpleSatsSetting.query.filter_by(school_id=pupil.school_id, academic_year=academic_year, exam_number=exam_number).first()
    default_maxima = {'arithmetic_score':40, 'reasoning_1_score':35, 'reasoning_2_score':35, 'reading_score':50, 'spelling_score':20, 'grammar_score':50}
    for key, value in values.items():
        if value is None:
            continue
        if 'scaled' in key:
            if not 80 <= value <= 120:
                raise CsvImportError(f'{key.replace("_", " ").title()} must be between 80 and 120.')
        elif key in default_maxima:
            maximum = getattr(configured, key.replace('_score','_max')) if configured else default_maxima[key]
            if not 0 <= value <= maximum:
                raise CsvImportError(f'{key.replace("_", " ").title()} must be between 0 and {maximum}.')
    matches = SatsResult.query.filter_by(school_id=pupil.school_id, pupil_id=pupil.id, academic_year=academic_year, exam_number=exam_number).all()
    if len(matches) > 1:
        raise CsvImportError('More than one saved SATs result matches this assessment point. Ask an admin to review it first.')
    result = matches[0] if matches else SatsResult(school_id=pupil.school_id, pupil_id=pupil.id, academic_year=academic_year, exam_number=exam_number, subject='maths', assessment_point=exam_number)
    for key, value in values.items():
        if value is not None:
            setattr(result,key,value)
    maths = (result.arithmetic_score,result.reasoning_1_score,result.reasoning_2_score)
    spag = (result.spelling_score,result.grammar_score)
    result.maths_combined_score = sum(maths) if all(value is not None for value in maths) else None
    result.spag_combined_score = sum(spag) if all(value is not None for value in spag) else None
    db.session.add(result)
    return result, not matches


def process_sats_workbook(ws, payload):
    from app.models import SimpleSatsExamTab
    report = {'errors':[], 'rows':[], 'added':0, 'changed_results':0}
    header_index = next(index for index in range(1,min(ws.max_row,5)+1) if any(str(cell.value or '').strip().lower() in {'pupil','pupil_id'} for cell in ws[index]))
    headers = [str(cell.value or '').strip().lower().replace(' ','_') for cell in ws[header_index]]
    aliases = {'arithmetic_score':('arithmetic','maths_arithmetic_raw'), 'reasoning_1_score':('reasoning_1','maths_reasoning_raw'), 'reasoning_2_score':('reasoning_2',), 'reading_score':('reading_paper','reading','reading_raw'), 'spelling_score':('spelling','spag_spelling_raw'), 'grammar_score':('grammar','spag_grammar_raw'), 'maths_scaled_score':('maths_scaled_score','maths_scaled'), 'reading_scaled_score':('reading_scaled','reading_scaled_score'), 'spag_scaled_score':('spag_scaled','spag_scaled_score')}
    for index, values in enumerate(ws.iter_rows(min_row=header_index+1,values_only=True),header_index+1):
        row = dict(zip(headers,values))
        try:
            scores = {field:parse_score(next((row[key] for key in names if row.get(key) not in (None,'')),None),field.replace('_',' ')) for field,names in aliases.items()}
            if not any(value is not None for value in scores.values()):
                continue
            pupil_id = parse_score(row.get('pupil_id'),'Pupil ID')
            pupil = Pupil.query.filter_by(id=pupil_id,school_id=payload['school_id']).first() if pupil_id else None
            if not pupil:
                candidates = Pupil.query.join(Pupil.school_class).filter(Pupil.school_id==payload['school_id'],SchoolClass.name==str(row.get('class') or '')).all()
                candidates = [pupil for pupil in candidates if pupil.full_name.casefold()==str(row.get('pupil') or '').casefold()]
                if len(candidates)!=1:
                    raise CsvImportError('Pupil could not be matched uniquely in this school.')
                pupil=candidates[0]
            point = str(row.get('assessment_point') or row.get('exam_number') or '').strip()
            fixed_points = {'Autumn 1':1,'Autumn 2':2,'Spring 1':3,'Spring 2':4}
            if point in fixed_points:
                exam = fixed_points[point]
            elif re.fullmatch(r'(?:Exam\s*)?[1-4]', point, re.IGNORECASE):
                exam = int(point[-1])
            else:
                tab = SimpleSatsExamTab.query.filter_by(school_id=payload['school_id'],academic_year=payload['academic_year'],name=point).first()
                if not tab:
                    raise CsvImportError('Assessment point must be Exam 1–4 or a configured SATs tab name.')
                exam=tab.exam_number
            result, created = save_sats_result(pupil,payload['academic_year'],exam,scores)
            report['added' if created else 'changed_results']+=1
            report['rows'].append({'pupil':pupil.full_name,'detail':f'SATs Exam {exam}: scores validated; scaled scores kept separate.', 'status':'ready'})
        except (ValueError,CsvImportError) as exc:
            report['errors'].append(f'SATs row {index}: {exc}')
    return report
