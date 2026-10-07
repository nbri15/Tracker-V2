"""Phase 2 boundaries, rollback, history and actual HTTP workflow tests."""
import base64
import io
from types import SimpleNamespace
import pytest
from flask_login import login_user
from sqlalchemy import text
from openpyxl import Workbook
from app import create_app
from app.extensions import db
from app.models import AcademicYear, AssessmentConfiguration, AssessmentReview, AssessmentSetting, AuditLog, GapQuestion, GapScore, GapTemplate, Intervention, Pupil, PupilClassHistory, School, SchoolClass, SubjectResult, User
from app.services.assessments import AssessmentValidationError, CsvImportError, compute_subject_result_values, get_setting_defaults, get_subject_setting, validate_setting_payload, _counts_from_band_labels
from app.services.assessment_reliability import apply_result, assessment_health, setup_impact, school_fingerprint
from app.services.assessment_imports import parse_csv_bytes, parse_score, simulate_import
from config import Config, config_by_name


@pytest.fixture
def world(tmp_path, monkeypatch):
    class TestConfig(Config):
        TESTING = True
        SECRET_KEY = 'phase-2-tests'
        SQLALCHEMY_DATABASE_URI = f'sqlite:///{tmp_path / "phase2.db"}'
        SQLALCHEMY_ENGINE_OPTIONS = {}
        WTF_CSRF_ENABLED = False
        DEMO_MODE = False
    monkeypatch.setitem(config_by_name, 'phase2test', TestConfig)
    app = create_app('phase2test')
    with app.app_context():
        db.create_all()
        year = AcademicYear(name='2026/27')
        old = AcademicYear(name='2025/26')
        a = School(name='School A', slug='school-a', current_academic_year=year)
        b = School(name='School B', slug='school-b', current_academic_year=year)
        teacher = User(username='teacher-a', role='teacher', school=a, is_active=True)
        admin = User(username='admin-a', role='school_admin', school=a, is_active=True)
        other = User(username='teacher-b', role='teacher', school=b, is_active=True)
        for user in (teacher, admin, other): user.set_password('password123')
        ca = SchoolClass(name='Oak', year_group=5, school=a, teacher=teacher, is_active=True)
        cb = SchoolClass(name='Oak', year_group=5, school=b, teacher=other, is_active=True)
        pupils = [Pupil(first_name=first, last_name='Test', gender='Female', school_class=ca, school_id=None, is_active=True) for first in ('Alice', 'Beth')]
        other_pupil = Pupil(first_name='Other', last_name='Pupil', gender='Male', school_class=cb, is_active=True)
        db.session.add_all([year, old, a, b, teacher, admin, other, ca, cb, *pupils, other_pupil]); db.session.flush()
        for pupil in pupils: pupil.school_id = a.id
        other_pupil.school_id = b.id
        db.session.commit()
        ids = dict(a=a.id, b=b.id, teacher=teacher.id, admin=admin.id, other=other.id, ca=ca.id, cb=cb.id, pupils=[pupil.id for pupil in pupils], other_pupil=other_pupil.id)
    yield app, ids


def login(client, user_id):
    with client.session_transaction() as session:
        session['_user_id'] = str(user_id); session['_fresh'] = True


def setting(subject='maths', threshold=55):
    return SimpleNamespace(school_id=None, **(get_setting_defaults(subject) | {'below_are_threshold_percent': threshold, 'on_track_threshold_percent': threshold}))


@pytest.mark.parametrize('subject,first,second,total', [('maths', 40,35,75),('reading',30,20,50),('spag',20,30,50)])
def test_subject_structures_and_maximum(subject, first, second, total):
    result = compute_subject_result_values(setting(subject), first, second)
    assert result == {'combined_score': total, 'combined_percent': 100.0, 'band_label': 'Exceeding'}


@pytest.mark.parametrize('first,second,total,band', [(0,0,0,'Working Towards'),(30,11,41,'Working Towards'),(30,12,42,'On Track'),(40,20,60,'Exceeding'),(None,20,None,None),(None,None,None,None)])
def test_zero_blank_and_boundaries(first,second,total,band):
    result = compute_subject_result_values(setting(), first, second)
    assert result['combined_score'] == total
    assert result['band_label'] == band


def test_exact_at_boundary_and_display_rounding():
    assert compute_subject_result_values(setting('reading', 60), 20,10)['band_label'] == 'On Track'
    assert SubjectResult.calculate_percent(1, 16) == 6.3


@pytest.mark.parametrize('score', [41,-1,1.5,float('nan'),float('inf')])
def test_invalid_scores(score):
    with pytest.raises(AssessmentValidationError): compute_subject_result_values(setting(), score,0)


@pytest.mark.parametrize('changes', [{'paper_1_max':0},{'combined_max':100},{'below_are_threshold_percent':-1},{'exceeding_threshold_percent':101},{'exceeding_threshold_percent':40},{'below_are_threshold_percent':float('nan')},{'on_track_threshold_percent':56}])
def test_impossible_setup(changes):
    with pytest.raises(AssessmentValidationError): validate_setting_payload(get_setting_defaults('maths') | changes)


@pytest.mark.parametrize('data', [b'', b'pupil,pupil\nA,B',b'pupil,paper_1_score\nA,1,2',b'pupil,paper_1_score\nA',b'pupil\n\xff',b'pupil\n"unclosed'])
def test_malformed_csv(data):
    with pytest.raises(CsvImportError): parse_csv_bytes(data)


def subject_payload(ids, rows):
    return {'school_id':ids['a'], 'class_id':ids['ca'], 'academic_year':'2026/27', 'subject':'maths', 'term':'autumn', 'rows':rows}


def test_preview_is_read_only_and_confirm_is_atomic(world):
    app, ids = world; client = app.test_client(); login(client, ids['teacher'])
    raw = f'pupil_id,pupil,paper_1_score,paper_2_score\n{ids["pupils"][0]},Alice Test,40,35\n{ids["pupils"][1]},Beth Test,0,0\n'.encode()
    response = client.post('/teacher/maths/upload', data={'csv_file':(io.BytesIO(raw),'scores.csv'), 'academic_year':'2026/27','term':'autumn'})
    assert response.status_code == 302
    with app.app_context(): assert SubjectResult.query.count() == 0
    preview = client.get(response.location)
    assert preview.status_code == 200 and b'2 matched' in preview.data
    saved = client.post(response.location, data={'action':'apply'}, follow_redirects=True)
    assert saved.status_code == 200 and b'0 failed' in saved.data and b'Complete' in saved.data
    with app.app_context():
        assert SubjectResult.query.count() == 2
        assert all(row.configuration_snapshot for row in SubjectResult.query.all())
        assert AuditLog.query.filter_by(action='assessment_results_imported').count() == 1
    again = client.post(response.location, data={'action':'apply'}, follow_redirects=True)
    assert b'already been applied' in again.data


@pytest.mark.parametrize('rows,expected', [
    ([{'pupil':'Alice Test','paper_1_score':'41','paper_2_score':'0'}], 'cannot exceed'),
    ([{'pupil':'Unknown','paper_1_score':'1','paper_2_score':'1'}], 'matched'),
    ([{'pupil':'Alice Test','paper_1_score':'1','paper_2_score':'1'}]*2, 'more than once'),
    ([{'pupil':'Alice Test','paper_1_score':'-1','paper_2_score':'1'}], 'below'),
    ([{'pupil':'Alice Test','paper_1_score':'oops','paper_2_score':'1'}], 'whole number'),
])
def test_invalid_imports_never_leave_partial_rows(world, rows, expected):
    app, ids = world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User, ids['teacher']))
        payload=subject_payload(ids, [{'pupil':'Beth Test','paper_1_score':'1','paper_2_score':'1'},*rows])
        report=simulate_import('subject',payload)
        assert any(expected in error for error in report['errors'])
        assert SubjectResult.query.count() == 0


def test_foreign_pupil_id_is_blocked(world):
    app, ids = world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['teacher']))
        report=simulate_import('subject', subject_payload(ids,[{'pupil_id':str(ids['other_pupil']), 'paper_1_score':'1','paper_2_score':'1'}]))
        assert report['errors'] and SubjectResult.query.count()==0


def test_threshold_preview_confirm_preserves_history_and_lower_year(world):
    app, ids=world; client=app.test_client(); login(client,ids['admin'])
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['admin']))
        pupil=db.session.get(Pupil,ids['pupils'][0]); pupil2=db.session.get(Pupil,ids['pupils'][1])
        for year in ('2025/26','2026/27'):
            result=SubjectResult(pupil=pupil,pupil_id=pupil.id,school_id=ids['a'],academic_year=year,term='autumn',subject='maths')
            apply_result(result, SimpleNamespace(school_id=ids['a'], **get_setting_defaults('maths')),30,10,cohort=5,assessment_year_group=5)
        low=SubjectResult(pupil=pupil2,pupil_id=pupil2.id,school_id=ids['a'],academic_year='2026/27',term='autumn',subject='maths')
        apply_result(low, SimpleNamespace(school_id=ids['a'], **get_setting_defaults('maths')),40,35,cohort=5,assessment_year_group=3)
        db.session.commit()
    response=client.post('/admin/assessments/setup',data={'academic_year':'2026/27','year_group':'5','subject':'maths','term':'autumn','paper_1_name':'Arithmetic','paper_1_max':'40','paper_2_name':'Reasoning','paper_2_max':'35','below_are_threshold_percent':'60','exceeding_threshold_percent':'80'})
    assert response.status_code==302
    page=client.get(response.location); assert b'1</strong> would change band' in page.data
    confirmed=client.post(response.location,data={'action':'apply'},follow_redirects=True)
    assert confirmed.status_code==200 and b'recalculated for 2026/27 only' in confirmed.data
    with app.app_context():
        rows=SubjectResult.query.filter_by(pupil_id=ids['pupils'][0]).all(); by_year={row.academic_year:row for row in rows}
        assert by_year['2025/26'].band_label=='On Track'
        assert by_year['2026/27'].band_label=='Working Towards'
        assert SubjectResult.query.filter_by(pupil_id=ids['pupils'][1]).one().band_label=='Working Towards'
        assert AssessmentConfiguration.query.count()==1
        assert AuditLog.query.filter_by(action='assessment_threshold_changed').count()==1


def test_legacy_bands_do_not_change_on_view_or_promotion(world):
    app, ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['teacher']))
        pupil=db.session.get(Pupil,ids['pupils'][0])
        result=SubjectResult(pupil=pupil,pupil_id=pupil.id,school_id=ids['a'],academic_year='2025/26',term='autumn',subject='maths',paper_1_score=30,paper_2_score=10,combined_score=40,combined_percent=53.3,band_label='On Track',assessment_year_group=5)
        db.session.add(result); db.session.commit()
        pupil.school_class.year_group=6; db.session.commit()
        assert _counts_from_band_labels([result])['On Track']==1
        assert result.configuration_snapshot is None


def test_manual_save_zero_blank_max_and_friendly_error(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    with app.app_context(): baseline=school_fingerprint(ids['a'])
    pupil=ids['pupils'][0]
    data={'form_name':'results','academic_year':'2026/27','term':'autumn','baseline':baseline,f'paper_1_score_{pupil}':'41', f'paper_2_score_{pupil}':'0',f'assessment_year_group_{pupil}':'5'}
    failed=client.post('/teacher/maths',data=data)
    assert failed.status_code==200 and b'cannot exceed 40' in failed.data and b'value="41"' in failed.data
    with app.app_context(): assert SubjectResult.query.count()==0
    data[f'paper_1_score_{pupil}']='0'
    response=client.post('/teacher/maths',data=data,follow_redirects=True)
    assert response.status_code==200 and b'All changes saved' in response.data
    with app.app_context():
        assert SubjectResult.query.one().combined_score==0
        data['baseline']=school_fingerprint(ids['a'])
    data[f'paper_1_score_{pupil}']=''; data[f'paper_2_score_{pupil}']=''
    response=client.post('/teacher/maths',data=data)
    assert b'clear scores' in response.data
    with app.app_context(): assert SubjectResult.query.one().combined_score==0


def test_stale_preview_and_wrong_owner(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    raw=f'pupil_id,paper_1_score,paper_2_score\n{ids["pupils"][0]},1,1'.encode()
    response=client.post('/teacher/maths/upload',data={'csv_file':(io.BytesIO(raw),'scores.csv')})
    with app.app_context():
        db.session.get(Pupil,ids['pupils'][0]).pupil_premium=True; db.session.commit()
    assert b'changed after this preview' in client.post(response.location,data={'action':'apply'}).data
    login(client,ids['other']); assert client.get(response.location).status_code==404
    with app.app_context(): assert SubjectResult.query.count()==0


def test_permissions_and_disabled_settings_autosave(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    assert client.get('/admin/assessments/setup').status_code==403
    login(client,ids['admin'])
    assert client.post('/admin/api/settings/quick-save',json={'field':'below_are_threshold_percent','value':60}).status_code==409
    assert client.get('/admin/settings').status_code==200


def test_workbook_calculates_all_years_and_rejects_invalid_scores(world):
    app,ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['admin']))
        wb=Workbook(); ws=wb.active; ws.title='Maths'
        ws.append(['pupil_id','pupil','class','term','arithmetic','reasoning','notes'])
        ws.append([ids['pupils'][0],'Alice Test','Oak','autumn',40,35,''])
        buffer=io.BytesIO(); wb.save(buffer)
        payload={'school_id':ids['a'],'academic_year':'2026/27','workbook':base64.b64encode(buffer.getvalue()).decode()}
        report=simulate_import('workbook',payload)
        assert report['errors']==[], report
        assert SubjectResult.query.count()==0
        from app.services.assessment_imports import process_import
        report=process_import('workbook',payload); db.session.commit()
        assert SubjectResult.query.one().combined_percent==100
        ws.cell(2,5).value=41; buffer=io.BytesIO(); wb.save(buffer); payload['workbook']=base64.b64encode(buffer.getvalue()).decode()
        report=simulate_import('workbook',payload)
        assert report['errors'] and SubjectResult.query.one().paper_1_score==40


def test_confirm_database_failure_rolls_everything_back(world, monkeypatch):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    raw=f'pupil_id,paper_1_score,paper_2_score\n{ids["pupils"][0]},1,1\n{ids["pupils"][1]},2,2'.encode()
    response=client.post('/teacher/maths/upload',data={'csv_file':(io.BytesIO(raw),'scores.csv')})
    from sqlalchemy.exc import OperationalError
    original=db.session.commit
    def fail(): raise OperationalError('simulated commit', {}, Exception('unavailable'))
    monkeypatch.setattr(db.session,'commit',fail)
    page=client.post(response.location,data={'action':'apply'})
    assert page.status_code==200 and b'No pupil results were saved' in page.data
    monkeypatch.setattr(db.session,'commit',original)
    with app.app_context():
        assert SubjectResult.query.count()==0
        assert AuditLog.query.count()==0
        assert AssessmentReview.query.one().confirmed_at is None


def test_school_settings_are_isolated(world):
    app,ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['teacher']))
        db.session.add(AssessmentSetting(school_id=ids['b'],year_group=5,subject='maths',term='autumn',**(get_setting_defaults('maths')|{'below_are_threshold_percent':99}))); db.session.commit()
        assert get_subject_setting(5,'maths','autumn','2026/27').below_are_threshold_percent==45
        assert AssessmentSetting.query.count()==1


def test_health_states(world):
    app,ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['teacher'])); pupil=db.session.get(Pupil,ids['pupils'][0]); configured=get_subject_setting(5,'maths','autumn','2026/27')
        assert assessment_health([pupil],{},configured)['label']=='Draft'
        configured.combined_max=0
        assert assessment_health([pupil],{},configured)['label']=='Configuration problem'


def test_sats_scaled_outcomes_are_separate(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    with app.app_context(): db.session.get(SchoolClass,ids['ca']).year_group=6; db.session.commit()
    page=client.get('/teacher/sats',follow_redirects=True)
    assert page.status_code==200
    from app.models import SimpleSatsSetting
    with app.app_context():
        setting=SimpleSatsSetting.query.first()
        assert setting is not None
        # Normal assessment configuration never changes SATs scale boundaries.
        assert AssessmentConfiguration.query.count()==0


def test_server_calculation_endpoint_scopes_and_validates(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    payload={'academic_year':'2026/27','term':'autumn','rows':[{'pupil_id':ids['pupils'][0],'paper_1_score':'40','paper_2_score':'35','assessment_year_group':3}]}
    page=client.post('/teacher/maths/calculate',json=payload)
    assert page.status_code==200 and page.json['rows'][0]['band_label']=='Working Towards'
    payload['rows'][0]['pupil_id']=ids['other_pupil']
    assert client.post('/teacher/maths/calculate',json=payload).status_code==403


@pytest.mark.parametrize('score,expected', [(99,'scaled-low'),(100,'scaled-at'),(110,'scaled-at'),(111,'scaled-high'),(None,'')])
def test_sats_scaled_boundaries(score,expected):
    from app.dashboards.routes import _scaled_band
    assert _scaled_band(score)==expected


def test_combined_csv_previews_without_creating_pupils_and_blocks_errors(world):
    app,ids=world; client=app.test_client(); login(client,ids['admin'])
    raw=b'pupil,class_name,year_group,maths_autumn_paper1,maths_autumn_paper2\nAlice Test,Oak,5,40,35\nNew Pupil,Oak,5,41,0'
    response=client.post('/admin/imports',data={'csv_file':(io.BytesIO(raw),'combined.csv'),'import_type':'combined','confirm_save':'1'})
    assert response.status_code==302
    preview=client.get(response.location)
    assert b'blocking problems' in preview.data and b'disabled' in preview.data
    client.post(response.location,data={'action':'apply'})
    with app.app_context():
        assert SubjectResult.query.count()==0
        assert Pupil.query.count()==3


def test_combined_csv_confirmation_preserves_pupil_flags(world):
    app,ids=world; client=app.test_client(); login(client,ids['admin'])
    with app.app_context():
        pupil=db.session.get(Pupil,ids['pupils'][0]); pupil.send=True; pupil.pupil_premium=True; db.session.commit()
    raw=b'pupil,class_name,year_group,maths_autumn_paper1,maths_autumn_paper2\nAlice Test,Oak,5,40,35'
    response=client.post('/admin/imports',data={'csv_file':(io.BytesIO(raw),'combined.csv'),'import_type':'combined'})
    page=client.post(response.location,data={'action':'apply'},follow_redirects=True)
    assert b'Import complete' in page.data
    with app.app_context():
        pupil=db.session.get(Pupil,ids['pupils'][0]); assert pupil.send and pupil.pupil_premium
        assert SubjectResult.query.one().combined_percent==100


def test_gap_sync_does_not_use_incomplete_or_fractional_totals(world):
    from app.services.gap import sync_gap_totals_to_subject_results
    app,ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['teacher'])); pupil=db.session.get(Pupil,ids['pupils'][0])
        template=GapTemplate(school_id=ids['a'],year_group=5,subject='maths',term='autumn',academic_year='2026/27')
        q1=GapQuestion(template=template,school_id=ids['a'],paper_key='paper_1',question_label='1',max_score=20)
        q2=GapQuestion(template=template,school_id=ids['a'],paper_key='paper_1',question_label='2',max_score=20)
        db.session.add_all([template,q1,q2]); db.session.flush()
        db.session.add(GapScore(pupil_id=pupil.id,question_id=q1.id,score=10,school_id=ids['a'])); db.session.flush()
        warnings=sync_gap_totals_to_subject_results([pupil],[q1,q2])
        assert warnings and SubjectResult.query.count()==0
        db.session.add(GapScore(pupil_id=pupil.id,question_id=q2.id,score=10.5,school_id=ids['a'])); db.session.flush()
        warnings=sync_gap_totals_to_subject_results([pupil],[q1,q2]); assert warnings and SubjectResult.query.count()==0
        GapScore.query.filter_by(question_id=q2.id).one().score=10; db.session.flush()
        assert not sync_gap_totals_to_subject_results([pupil],[q1,q2])
        assert SubjectResult.query.one().paper_1_score==20
        assert SubjectResult.query.one().combined_percent is None


def test_old_cohort_survives_manual_edit_after_promotion(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    with app.app_context():
        pupil=db.session.get(Pupil,ids['pupils'][0]); pupil.school_class.year_group=6
        db.session.add(PupilClassHistory(school_id=ids['a'],pupil_id=pupil.id,academic_year='2025/26',class_name='Oak',year_group=5)); db.session.commit()
        baseline=school_fingerprint(ids['a'])
    data={'form_name':'results','academic_year':'2025/26','term':'autumn','baseline':baseline,f'paper_1_score_{ids["pupils"][0]}':'40',f'paper_2_score_{ids["pupils"][0]}':'35',f'assessment_year_group_{ids["pupils"][0]}':'5'}
    response=client.post('/teacher/maths',data=data,follow_redirects=True)
    assert response.status_code==200
    with app.app_context():
        result=SubjectResult.query.one(); assert result.cohort_year_group==5 and result.band_label=='Exceeding'


def test_expired_review_cannot_import(world):
    from datetime import datetime, timedelta, timezone
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    raw=f'pupil_id,paper_1_score,paper_2_score\n{ids["pupils"][0]},1,1'.encode()
    response=client.post('/teacher/maths/upload',data={'csv_file':(io.BytesIO(raw),'scores.csv')})
    with app.app_context():
        AssessmentReview.query.one().created_at=datetime.now(timezone.utc)-timedelta(hours=1); db.session.commit()
    assert b'expired' in client.post(response.location,data={'action':'apply'}).data
    with app.app_context(): assert SubjectResult.query.count()==0


def test_maximum_reduction_blocks_setup_without_mutation(world):
    app,ids=world; client=app.test_client(); login(client,ids['admin'])
    with app.app_context():
        pupil=db.session.get(Pupil,ids['pupils'][0]); row=SubjectResult(pupil=pupil,pupil_id=pupil.id,academic_year='2026/27',term='autumn',subject='maths',school_id=ids['a'])
        apply_result(row,SimpleNamespace(school_id=ids['a'],**get_setting_defaults('maths')),40,35,cohort=5,assessment_year_group=5); db.session.commit()
    response=client.post('/admin/assessments/setup',data={'academic_year':'2026/27','year_group':'5','subject':'maths','term':'autumn','paper_1_name':'Arithmetic','paper_1_max':'20','paper_2_name':'Reasoning','paper_2_max':'35','below_are_threshold_percent':'55','exceeding_threshold_percent':'80'})
    preview=client.get(response.location); assert b'cannot exceed 20' in preview.data
    client.post(response.location,data={'action':'apply'})
    with app.app_context():
        assert AssessmentConfiguration.query.count()==0
        assert SubjectResult.query.one().combined_percent==100


def test_invalid_bulk_save_rolls_back_other_pupils(world):
    app,ids=world; client=app.test_client(); login(client,ids['teacher'])
    with app.app_context(): baseline=school_fingerprint(ids['a'])
    data={'form_name':'results','academic_year':'2026/27','term':'autumn','baseline':baseline}
    for pupil,first in zip(ids['pupils'],['20','41']):
        data.update({f'paper_1_score_{pupil}':first,f'paper_2_score_{pupil}':'10',f'assessment_year_group_{pupil}':'5'})
    response=client.post('/teacher/maths',data=data)
    assert b'cannot exceed' in response.data
    with app.app_context(): assert SubjectResult.query.count()==0


def test_csrf_is_required_for_confirmation(world):
    app,ids=world; app.config['WTF_CSRF_ENABLED']=True
    client=app.test_client(); login(client,ids['teacher'])
    assert client.post('/teacher/maths/upload',data={}).status_code==400


def test_migration_preserves_legacy_results_and_other_data(world):
    from flask_migrate import stamp, upgrade
    app,ids=world
    with app.app_context():
        pupil=ids['pupils'][0]
        db.session.add(SubjectResult(pupil_id=pupil,school_id=ids['a'],academic_year='2025/26',subject='maths',term='autumn',paper_1_score=30,paper_2_score=10,combined_score=40,combined_percent=53.3,band_label='On Track'))
        db.session.add(Intervention(pupil_id=pupil,school_id=ids['a'],subject='maths',term='autumn',academic_year='2025/26',reason='Historical support',is_active=True))
        db.session.add(AssessmentSetting(school_id=ids['a'],year_group=5,subject='maths',term='autumn',**get_setting_defaults('maths'))); db.session.commit()
        # Isolated test DB only: recreate the actual pre-phase-2 schema.
        db.session.execute(text('DROP TABLE assessment_reviews'))
        db.session.execute(text('DROP TABLE assessment_configurations'))
        db.session.execute(text('ALTER TABLE subject_results DROP COLUMN configuration_snapshot'))
        db.session.execute(text('ALTER TABLE subject_results DROP COLUMN cohort_year_group')); db.session.commit()
        stamp(revision='20260922_01'); upgrade()
        result=SubjectResult.query.one()
        assert result.band_label=='On Track' and result.combined_percent==53.3
        assert result.configuration_snapshot is None
        assert Intervention.query.one().reason=='Historical support'
        assert AssessmentSetting.query.one().below_are_threshold_percent==45
        assert Pupil.query.count()==3


def test_downloaded_full_workbook_roundtrip_keeps_flags_and_join_year(world):
    from app.admin.routes import _build_full_template_workbook
    from app.services.assessment_imports import process_import
    app,ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['admin'])); pupil=db.session.get(Pupil,ids['pupils'][0]); pupil.join_year_group=0; pupil.send=True; db.session.commit()
        wb=_build_full_template_workbook(ids['a']); buffer=io.BytesIO(); wb.save(buffer)
        payload={'school_id':ids['a'],'academic_year':'2026/27','workbook':base64.b64encode(buffer.getvalue()).decode()}
        report=simulate_import('workbook',payload)
        assert not report['errors'],report['errors']
        process_import('workbook',payload); db.session.commit()
        assert db.session.get(Pupil,pupil.id).join_year_group==0
        assert db.session.get(Pupil,pupil.id).send is True
        assert SubjectResult.query.count()==0


def test_sats_workbook_feeds_active_tracker_and_keeps_partial_totals_missing(world):
    from app.services.assessment_imports import process_import
    from app.models import SatsResult
    app,ids=world
    with app.app_context(), app.test_request_context('/'):
        login_user(db.session.get(User,ids['admin'])); db.session.get(SchoolClass,ids['ca']).year_group=6; db.session.commit()
        wb=Workbook(); ws=wb.active; ws.title='SATs'
        ws.append(['pupil_id','Pupil','Class','Year Group','Assessment Point','Arithmetic','Reasoning 1','Reasoning 2','Maths Scaled Score','Reading Paper','Reading Scaled','Spelling','Grammar','SPaG Scaled'])
        ws.append([ids['pupils'][0],'Alice Test','Oak',6,'Autumn 1',40,35,None,100,50,99,20,50,110])
        buffer=io.BytesIO(); wb.save(buffer)
        payload={'school_id':ids['a'],'academic_year':'2026/27','workbook':base64.b64encode(buffer.getvalue()).decode()}
        report=simulate_import('workbook',payload); assert not report['errors'],report
        assert SatsResult.query.count()==0
        report=process_import('workbook',payload); db.session.commit()
        result=SatsResult.query.one(); assert result.maths_scaled_score==100 and result.reading_scaled_score==99
        assert result.maths_combined_score is None and result.spag_combined_score==70
        assert SubjectResult.query.count()==0
